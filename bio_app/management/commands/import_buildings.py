import openpyxl
from django.core.management.base import BaseCommand
from bio_app.models import VacantBuilding
from datetime import date


# Mapping from xlsx values to model choice keys
TYPE_OF_USE_MAP = {
    'Temporary': 'temporary',
    'Long-term': 'long_term',
    'Permanent': 'permanent',
}
PREVIOUS_USE_MAP = {
    'Residential': 'residential',
    'Industry': 'industry',
    'Gastronomy': 'gastronomy',
    'Small retail space': 'small_retail',
    'Large retail space': 'large_retail',
}
FACILITY_SIZE_MAP = {
    '15–40 m²': '15_40',
    '40–100 m²': '40_100',
    '100–200 m²': '100_200',
    '>200 m²': 'over_200',
}
REACHABILITY_MAP = {
    'Very Central': 'very_central',
    'Central': 'central',
    'Somewhat Central': 'somewhat_central',
    'not centraly located': 'not_central',
}
GOVERNANCE_MAP = {
    'civil society organisations': 'civil_society',
    'public': 'public',
    'private': 'private',
    'mixed or other models': 'mixed',
}
PREVIOUS_STATE_MAP = {
    'poor': 'poor',
    'moderate': 'moderate',
    'excellent': 'excellent',
    'excelent': 'excellent',  # handle typo in xlsx
}

# Only import real named buildings (skip single-letter placeholders)
REAL_BUILDINGS = {
    'Altes Milchhaus',
    'Arbeit im Dorf',
    'Die Giesserei',
    'Leutkircher Bürgerbahnhof',
    'Dorfplatz St.Andrä-Wördern',
}


class Command(BaseCommand):
    help = 'Import vacant buildings from the xlsx spreadsheet'

    def add_arguments(self, parser):
        parser.add_argument(
            '--all', action='store_true',
            help='Import all rows including placeholder data (single-letter names)',
        )

    def handle(self, *args, **options):
        import_all = options['all']
        xlsx_path = 'requirements-gitignore/Vacant Buildings -20260216T162608Z-1-001/Vacant Buildings/Vacant Buildings_factsheet_categories.xlsx'

        wb = openpyxl.load_workbook(xlsx_path)
        ws = wb['Sheet2']

        created = 0
        skipped = 0
        for i, row in enumerate(ws.iter_rows(min_row=2, values_only=True)):
            name = str(row[0]).strip() if row[0] else ''
            if not name or name == 'all':
                continue

            if not import_all and name not in REAL_BUILDINGS:
                skipped += 1
                continue

            type_of_use = TYPE_OF_USE_MAP.get(str(row[1]).strip(), '') if row[1] else ''
            previous_use = PREVIOUS_USE_MAP.get(str(row[2]).strip(), '') if row[2] else ''
            facility_size = FACILITY_SIZE_MAP.get(str(row[3]).strip(), '') if row[3] else ''
            reachability = REACHABILITY_MAP.get(str(row[4]).strip(), '') if row[4] else ''
            governance = GOVERNANCE_MAP.get(str(row[5]).strip(), '') if row[5] else ''
            previous_state = PREVIOUS_STATE_MAP.get(str(row[6]).strip(), '') if row[6] else ''

            obj, was_created = VacantBuilding.objects.update_or_create(
                name=name,
                defaults={
                    'address': '',
                    'area_sq_ft': 0,
                    'available_from': date.today(),
                    'year': 2024,
                    'floor': 0,
                    'type_of_use': type_of_use or None,
                    'previous_use': previous_use or None,
                    'facility_size': facility_size or None,
                    'reachability': reachability or None,
                    'governance': governance or None,
                    'previous_state': previous_state or None,
                },
            )
            action = 'Created' if was_created else 'Updated'
            self.stdout.write(f'  {action}: {name}')
            created += 1

        self.stdout.write(self.style.SUCCESS(
            f'\nDone. Imported {created} buildings, skipped {skipped} placeholder rows.'
        ))
